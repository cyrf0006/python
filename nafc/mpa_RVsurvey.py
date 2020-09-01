'''
    Records are:
    Record #0: [5, 'Northeast Newfoundland Slope Closure 1', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 6, 1443556.74558, 54111.5220788]
    Record #1: [6, 'Northeast Newfoundland Slope Closure 2', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 5, 211242.748802, 1154.68404126]
    Record #2: [7, 'Hopedale Saddle Closure', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 2, 818222.095658, 15412.4572173]
    Record #3: [9, 'Funk Island Deep Closure', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 4, 378737.019016, 7274.21207849]
    Record #4: [14, 'Hawke Channel Closure', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 3, 376093.587497, 8836.57491554]
    Record #5: [15, 'Laurentian Channel', 'Marine Protected Area', 'Fisheries and Oceans Canada', 3, 678661.298122, 11619.4045159]
    Record #6: [35, 'Division 3O Coral and Sponge Conservation Closure', 'Marine Refuge (OEABCM)', 'Fisheries and Oceans Canada', 7, 679455.445634, 10422.0091543]
    Record #0: [0, '3O Coral Closure']
    Record #0: [1, 0, 'Corner Seamounts', 9.9999999998, 3.9999999993]
    Record #1: [2, 0, 'Fogo Seamounts 1', 2.84722222053, 0.493217591325]
    Record #2: [3, 0, 'Fogo Seamounts 2', 2.84666668052, 0.493052088991]
    Record #3: [4, 0, 'Newfoundland Seamounts', 7.70000002058, 1.72222223814]
    Record #4: [5, 0, 'Orphan Knoll', 5.9999999988, 1.9999999992]
    Record #5: [6, 0, 'New England Seamounts', 24.5133095399, 18.0128306383]
    Record #0: [1, 'Tail of the Bank']
    Record #1: [2, 'Flemish Pass / Eastern Canyon']
    Record #2: [3, 'Beothuk Knoll']
    Record #3: [4, 'Eastern Flemish Cap']
    Record #4: [5, 'Northeast Flemish Cap']
    Record #5: [6, 'Sackville Spur']
    Record #6: [7, 'Northern Flemish Cap']
    Record #7: [8, 'Northern Flemish Cap']
    Record #8: [9, 'Northern Flemish Cap']
    Record #9: [10, 'Northwest Flemish Cap']
    Record #10: [11, 'Northwest Flemish Cap']
    Record #11: [12, 'Northwest Flemish Cap']
    Record #12: [13, 'Beothuk Knoll']
    
'''
import netCDF4
import h5py                                                                
import os
os.environ['PROJ_LIB'] = '/home/cyrf0006/anaconda3/share/proj'
from mpl_toolkits.basemap import Basemap
import numpy as  np
import matplotlib.pyplot as plt
import cmocean
import cmocean.cm as cmo
import cartopy.crs as ccrs
from cartopy.mpl.gridliner import LONGITUDE_FORMATTER, LATITUDE_FORMATTER
import cartopy.feature as cpf
from cartopy.mpl.geoaxes import GeoAxes
import matplotlib.ticker as mticker
import openpyxl, pprint
import shapefile 
import pandas as pd
from shapely.geometry import Point
from shapely.geometry.polygon import Polygon
from bokeh.plotting import figure, save
#import geopandas as gpd

## ---- Region parameters ---- ##
dataFile = '/home/cyrf0006/data/GEBCO/GEBCO_2014_1D.nc'
lon_0 = -50
lat_0 = 50
lonLims = [-65, -40]
latLims = [40, 60]
proj = 'merc'
decim_scale = 4
fig_name = 'map_closures.png'
v = np.linspace(0, 3500, 36)

## ---- read shapefile ---- ##
# MPAS / DFO closures
sf = shapefile.Reader('/home/cyrf0006/research/MPAs/Warren_shapefiles/RV_CSAS_closures_GCS')
records = sf.records()
shapes = sf.shapes()
# Fill dictionary with closures (named MPAs for simplicity)
mpas = {}
for idx, rec in enumerate(records):
    if rec[0] == '':
        continue
    else:
        print(rec)
        mpas[rec[0]] = np.array(shapes[idx].points)

# NAFO closures
# 3O


        sf = shapefile.Reader('/home/cyrf0006/research/MPAs/NAFO_closures/2015_Closures_3O')
records = sf.records()
shapes = sf.shapes()
nafo_3O = {}
for idx, rec in enumerate(records):
    if rec[0] == '':
        continue
    else:
        print(rec)
        nafo_3O[rec[0]] = np.array(shapes[idx].points)
# Seamounts
sf = shapefile.Reader('/home/cyrf0006/research/MPAs/NAFO_closures/2018_Closures_seamounts')
records = sf.records()
shapes = sf.shapes()
nafo_seamounts = {}
for idx, rec in enumerate(records):
    if rec[0] == '':
        continue
    else:
        print(rec)
        nafo_seamounts[rec[0]] = np.array(shapes[idx].points)
# Sponge&Corals
sf = shapefile.Reader('/home/cyrf0006/research/MPAs/NAFO_closures/2019_Closures_sponge_coral')
records = sf.records()
shapes = sf.shapes()
nafo_coral = {}
for idx, rec in enumerate(records):
    if rec[0] == '':
        continue
    else:
        print(rec)
        nafo_coral[rec[0]] = np.array(shapes[idx].points)        

## ---- Bathymetry ---- ##
print('Load and grid bathymetry')
# h5 file
h5_outputfile = 'mpa_bathymetry.h5'
if os.path.isfile(h5_outputfile):
     print([h5_outputfile + ' exists! Reading directly'])
     h5f = h5py.File(h5_outputfile,'r')
     lon = h5f['lon'][:]
     lat = h5f['lat'][:]
     Z = h5f['Z'][:]
     h5f.close()

else:
    # Extract variables
    dataset = netCDF4.Dataset(dataFile)
    x = [-179-59.75/60, 179+59.75/60] # to correct bug in 30'' dataset?
    y = [-89-59.75/60, 89+59.75/60]
    spacing = dataset.variables['spacing']

    # Compute Lat/Lon
    nx = int((x[-1]-x[0])/spacing[0]) + 1  # num pts in x-dir
    ny = int((y[-1]-y[0])/spacing[1]) + 1  # num pts in y-dir
    lon = np.linspace(x[0],x[-1],nx)
    lat = np.linspace(y[0],y[-1],ny)
    # Reshape data
    zz = dataset.variables['z']
    Z = zz[:].reshape(ny, nx)
    Z = np.flipud(Z) # <------------ important!!!
    # Reduce data according to Region params
    idx_lon = np.where((lon>=lonLims[0]) & (lon<=lonLims[1]))
    idx_lat = np.where((lat>=latLims[0]) & (lat<=latLims[1]))
    Z = Z[idx_lat[0][0]:idx_lat[0][-1]+1, idx_lon[0][0]:idx_lon[0][-1]+1]
    lon = lon[idx_lon[0]]
    lat = lat[idx_lat[0]]

    # Save data for later use
    h5f = h5py.File(h5_outputfile, 'w')
    h5f.create_dataset('lon', data=lon)
    h5f.create_dataset('lat', data=lat)
    h5f.create_dataset('Z', data=Z)
    h5f.close()
    print(' -> Done!')

# Reduce data according to Region params
lon = lon[::decim_scale]
lat = lat[::decim_scale]
Z = Z[::decim_scale, ::decim_scale]

#plt.title('Standard air temperature sites')
print('--- Now plot ---')
fig = plt.figure(figsize=(7,5))
ax = fig.add_subplot(111, projection=ccrs.Mercator())
ax.set_extent([lonLims[0], lonLims[1], latLims[0], latLims[1]], crs=ccrs.PlateCarree())
ax.add_feature(cpf.NaturalEarthFeature('physical', 'coastline', '50m', edgecolor='k', alpha=0.7, linewidth=0.6, facecolor='black'), zorder=1)
m=ax.gridlines(linewidth=0.5, color='black', draw_labels=True, alpha=0.5)
m.xlabels_top=False
m.ylabels_right=False
m.xlocator = mticker.FixedLocator([-75, -70, -60, -50, -40])
m.ylocator = mticker.FixedLocator([40, 45, 50, 55, 60, 65])
m.xformatter = LONGITUDE_FORMATTER
m.yformatter = LATITUDE_FORMATTER
m.ylabel_style = {'size': 7, 'color': 'black', 'weight':'bold'}
m.xlabel_style = {'size': 7, 'color': 'black', 'weight':'bold'}
lightdeep = cmocean.tools.lighten(cmo.deep, 0.5)
ls = np.linspace(0, 5500, 20)
c = plt.contourf(lon, lat, -Z, ls, transform=ccrs.PlateCarree(), cmap=lightdeep, extend='max', zorder=5)
cc = plt.contour(lon, lat, -Z, [100, 500, 1000, 2000, 3000, 4000, 5000], colors='silver', linewidths=.5, transform=ccrs.PlateCarree(), zorder=10)
plt.clabel(cc, inline=True, fontsize=7, fmt='%i')

# plot MPAs/closures
# DFO's
for idx, key in enumerate(mpas.keys()):
    coords = mpas[key]  
    poly_x = coords[:,0]
    poly_y = coords[:,1]
    poly = list(zip(poly_x, poly_y))
    pgon = Polygon(poly)
    ax.add_geometries([pgon], crs=ccrs.PlateCarree(), facecolor='red', alpha=0.5)
    # add numer
    ax.text(poly_x.mean(), poly_y.mean(), str(records[idx][0]), transform=ccrs.PlateCarree(), horizontalalignment='center', verticalalignment='center', fontsize=10, color='r', fontweight='bold', zorder=10)
# NAFO's 3O
for idx, key in enumerate(nafo_3O.keys()):
    coords = nafo_3O[key]  
    poly_x = coords[:,0]
    poly_y = coords[:,1]
    poly = list(zip(poly_x, poly_y))
    pgon = Polygon(poly)
    ax.add_geometries([pgon], crs=ccrs.PlateCarree(), facecolor='orange', alpha=0.5)
    # add numer
    ax.text(poly_x.mean(), poly_y.mean(), str(records[idx][0]), transform=ccrs.PlateCarree(), horizontalalignment='center', verticalalignment='center', fontsize=10, color='r', fontweight='bold', zorder=10)
# NAFO's Seamounts
for idx, key in enumerate(nafo_seamounts.keys()):
    coords = nafo_seamounts[key]  
    poly_x = coords[:,0]
    poly_y = coords[:,1]
    poly = list(zip(poly_x, poly_y))
    pgon = Polygon(poly)
    ax.add_geometries([pgon], crs=ccrs.PlateCarree(), facecolor='orange', alpha=0.5)
    # add numer
    ax.text(poly_x.mean(), poly_y.mean(), str(records[idx][0]), transform=ccrs.PlateCarree(), horizontalalignment='center', verticalalignment='center', fontsize=10, color='r', fontweight='bold', zorder=10)
# NAFO's Sponge & Corals
for idx, key in enumerate(nafo_coral.keys()):
    coords = nafo_coral[key]  
    poly_x = coords[:,0]
    poly_y = coords[:,1]
    poly = list(zip(poly_x, poly_y))
    pgon = Polygon(poly)
    ax.add_geometries([pgon], crs=ccrs.PlateCarree(), facecolor='orange', alpha=0.5)
    # add numer
    ax.text(poly_x.mean(), poly_y.mean(), str(records[idx][0]), transform=ccrs.PlateCarree(), horizontalalignment='center', verticalalignment='center', fontsize=10, color='r', fontweight='bold', zorder=10)


    
#### ---- Save Figure ---- ####
fig.set_size_inches(w=10, h=12)
fig.savefig(fig_name, dpi=200)
os.system('convert -trim -bordercolor White -border 30x30 ' + fig_name + ' ' + fig_name)
#plt.show()

